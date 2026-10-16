"""Build the editable manuscript from MANUSCRIPT.md (requires python-docx)."""
from datetime import datetime, timezone
from pathlib import Path
import re
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT

HERE = Path(__file__).resolve().parent


def inline(paragraph, text):
    pattern = r'(\[[^\]]+\]\(https?://[^)]+\)|\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)'
    for part in re.split(pattern, text):
        match = re.fullmatch(r'\[([^\]]+)\]\((https?://[^)]+)\)', part)
        if match:
            link = OxmlElement('w:hyperlink')
            link.set(qn('r:id'), paragraph.part.relate_to(match[2], RT.HYPERLINK, is_external=True))
            run = OxmlElement('w:r'); props = OxmlElement('w:rPr')
            color = OxmlElement('w:color'); color.set(qn('w:val'), '174A72'); props.append(color)
            run.append(props); node = OxmlElement('w:t'); node.text = match[1]; run.append(node)
            link.append(run); paragraph._p.append(link)
        elif part.startswith('**'):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith('`'):
            run = paragraph.add_run(part[1:-1]); run.font.name = 'Courier New'; run.font.size = Pt(9)
        elif part.startswith('*'):
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def table(doc, rows, number):
    widths = {1: [2.0, 1.45, 3.05], 2: [2.7, 1.15, 1.15, 1.5],
              3: [2.2, 4.3], 4: [2.15, 4.35]}[number]
    grid = doc.add_table(rows=0, cols=len(rows[0])); grid.alignment = WD_TABLE_ALIGNMENT.CENTER
    grid.autofit = False
    for col, width in zip(grid.columns, widths): col.width = Inches(width)
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        node = OxmlElement('w:' + edge); node.set(qn('w:val'), 'single')
        node.set(qn('w:sz'), '4'); node.set(qn('w:color'), 'D0D0D0'); borders.append(node)
    grid._tbl.tblPr.append(borders)
    for index, values in enumerate(rows):
        row = grid.add_row()
        tr_pr = row._tr.get_or_add_trPr(); tr_pr.append(OxmlElement('w:cantSplit'))
        if index == 0:
            tr_pr.append(OxmlElement('w:tblHeader'))
        for cell, value, width in zip(row.cells, values, widths):
            cell.width = Inches(width); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(5); p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.line_spacing = 1.04
            inline(p, value)
            for r in p.runs: r.font.size = Pt(10); r.bold = index == 0
            if index == 0:
                shading = OxmlElement('w:shd'); shading.set(qn('w:fill'), 'F2F2F2'); cell._tc.get_or_add_tcPr().append(shading)
    return grid


def main():
    doc = Document(); sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
    sec.header_distance = sec.footer_distance = Inches(.5)
    for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
        s = doc.styles[name]; s.font.name = 'Times New Roman'; s.font.color.rgb = RGBColor(0, 0, 0)
        fonts = s.element.get_or_add_rPr().get_or_add_rFonts()
        for key in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
            fonts.attrib.pop(qn('w:' + key), None)
        for key in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
            fonts.set(qn('w:' + key), 'Times New Roman')
        for border in s.element.findall('.//' + qn('w:pBdr')):
            border.getparent().remove(border)
    normal = doc.styles['Normal']; normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.12; normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    for name, size, before in [('Heading 1', 13, 12), ('Heading 2', 11.5, 9)]:
        s = doc.styles[name]; s.font.size = Pt(size); s.font.bold = True
        s.paragraph_format.space_before = Pt(before); s.paragraph_format.space_after = Pt(5)
        s.paragraph_format.keep_with_next = True
    doc.styles['Title'].font.size = Pt(20)
    doc.styles['Title'].paragraph_format.space_after = Pt(9)
    header = sec.header.paragraphs[0]; header.text = 'VERIFYING STOP SCOPE IN AGENT WORKFLOWS'
    header.runs[0].font.size = Pt(8); header.runs[0].font.color.rgb = RGBColor.from_string('666666')
    footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('Buldurgan · v1.2 · ').font.size = Pt(9)
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    props = doc.core_properties
    props.title = 'Verifying Stop Scope in Agent Workflows'
    props.author = 'Huseyin Buldurgan'; props.subject = 'Accepted work, restart, and output boundaries'
    props.version = '1.2'; props.keywords = 'agent workflows; cancellation; revocation; reproducibility; AI safety'
    props.created = props.modified = datetime(2026, 10, 16, tzinfo=timezone.utc)
    blocks = (HERE / 'MANUSCRIPT.md').read_text().strip().split('\n\n')
    table_number, in_refs, after_table = 0, False, False
    for block in blocks:
        if block.startswith('# '):
            p = doc.add_paragraph(block[2:], 'Title'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif block.startswith('## '):
            in_refs = block == '## References'
            doc.add_paragraph(block[3:], 'Heading 1')
        elif block.startswith('### '):
            doc.add_paragraph(block[4:], 'Heading 2')
        elif block.startswith('|'):
            rows = [[c.strip() for c in line.strip().strip('|').split('|')] for line in block.splitlines()
                    if not re.fullmatch(r'[| :\-]+', line)]
            table_number += 1; table(doc, rows, table_number); after_table = True
        else:
            p = doc.add_paragraph(); inline(p, ' '.join(block.splitlines()))
            if after_table:
                p.paragraph_format.space_before = Pt(8); after_table = False
            if block.startswith('**Table '):
                p.paragraph_format.keep_with_next = True
                for r in p.runs: r.font.size = Pt(10)
            elif block.startswith(('Huseyin Buldurgan', 'Independent Researcher', 'Research preprint')):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(4)
            elif in_refs:
                p.paragraph_format.left_indent = Inches(.22); p.paragraph_format.first_line_indent = Inches(-.22)
                for r in p.runs: r.font.size = Pt(10)
    output = HERE / 'Verifying_Stop_Scope_v1.2.docx'; doc.save(output)
    print(output.name)


if __name__ == '__main__':
    main()
